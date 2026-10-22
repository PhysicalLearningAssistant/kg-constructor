"""教材读取：Markdown / TXT / DOCX / PDF → 带页码标记的纯文本。"""

from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass
from pathlib import Path

PAGE_MARKER = "[[PAGE:{page}]]"


@dataclass
class LoadedText:
    path: Path
    text: str
    source_type: str
    sha256: str
    page_count: int | None
    title: str


def _decode(raw: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", "replace")


def _docx_to_markdown(path: Path) -> str:
    try:
        import docx  # type: ignore
    except ModuleNotFoundError as exc:  # pragma: no cover - 取决于环境
        raise RuntimeError("读取 .docx 需要安装 python-docx：pip install python-docx") from exc

    document = docx.Document(str(path))
    lines: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style_name = (paragraph.style.name or "") if paragraph.style else ""
        match = re.search(r"(?:Heading|标题)\s*(\d+)", style_name, re.I)
        if match:
            level = max(1, min(6, int(match.group(1))))
            lines.append(f"{'#' * level} {text}")
        else:
            lines.append(text)
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))
    return "\n\n".join(lines)


def _pdf_to_text(path: Path) -> tuple[str, int]:
    try:
        from pypdf import PdfReader  # type: ignore
    except ModuleNotFoundError as exc:  # pragma: no cover - 取决于环境
        raise RuntimeError("读取 .pdf 需要安装 pypdf：pip install pypdf") from exc

    reader = PdfReader(str(path))
    pages: list[str] = []
    for index, page in enumerate(reader.pages, start=1):
        pages.append(PAGE_MARKER.format(page=index) + "\n" + (page.extract_text() or ""))
    return "\n\n".join(pages), len(reader.pages)


def read_textbook(path: Path, title: str | None = None) -> LoadedText:
    raw = path.read_bytes()
    suffix = path.suffix.lower()
    page_count: int | None = None

    if suffix == ".docx":
        text = _docx_to_markdown(path)
        source_type = "docx"
    elif suffix == ".pdf":
        text, page_count = _pdf_to_text(path)
        source_type = "pdf"
    elif suffix in {".md", ".markdown", ".txt"}:
        text = _decode(raw)
        source_type = "markdown"
    else:
        raise ValueError(f"不支持的教材格式：{suffix}（支持 .md/.txt/.docx/.pdf）")

    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if not text:
        raise ValueError(f"教材内容为空：{path}")

    return LoadedText(
        path=path,
        text=text,
        source_type=source_type,
        sha256=hashlib.sha256(raw).hexdigest(),
        page_count=page_count,
        title=title or path.stem,
    )
